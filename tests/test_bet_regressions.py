"""Synthetic regressions for BET recalculation, units, and exported results."""
import copy
import io
import os
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import Workbook, load_workbook

from modules import bet_processor as bet


def synthetic_isotherm():
    pressure = np.linspace(0.01, 0.4, 27)
    qm, c = 6.0, 85.0
    # Slight curvature makes changing the window a meaningful regression.
    quantity = qm * c * pressure / ((1 - pressure) * (1 + (c - 1) * pressure))
    quantity *= 1 + 0.7 * pressure ** 2
    return {
        'sample_id': 'Synthetic BET', 'sample_mass_g': 0.1, 'adsorptive': 'N2',
        'quantity_basis': 'per_mass', 'quantity_header': 'Quantity Adsorbed (cm³/g STP)',
        'source_file': 'synthetic.xlsx', 'source_metrics': {'BET surface area': '24.0 m²/g'},
        'points': [
            {'branch': 'Adsorption', 'relative_pressure': float(x),
             'quantity_cm3_g_stp': float(q), 'source_quantity': float(q)}
            for x, q in zip(pressure, quantity)
        ],
    }


def synthetic_workbook(mass='0.1000 g', header='Quantity Adsorbed (cm³/g STP)'):
    data = synthetic_isotherm()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(('Sample:', 'Synthetic BET'))
    if mass is not None:
        sheet.append(('Sample Mass:', mass))
    sheet.append(('BET surface area:', '24.0 m²/g'))
    sheet.append(('Relative Pressure (p/p₀)', header))
    for point in data['points']:
        sheet.append((point['relative_pressure'], point['source_quantity']))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class BetRecalculationTests(unittest.TestCase):
    def test_mass_changes_area_and_volume_without_mutating_source_or_shape(self):
        parsed = synthetic_isotherm()
        original = copy.deepcopy(parsed)
        baseline = bet.analyze_bet(parsed, .05, .3)
        for mass in (.05, .2, .08):
            with self.subTest(mass=mass):
                result = bet.analyze_bet(parsed, .05, .3, sample_mass_g=mass)
                factor = .1 / mass
                for key in ('surface_area_m2_g', 'monolayer_capacity_cm3_g_stp',
                            'total_pore_volume_cm3_g'):
                    self.assertAlmostEqual(result[key], baseline[key] * factor, places=10)
                for key in ('c_constant', 'average_pore_diameter_nm', 'r_squared'):
                    self.assertAlmostEqual(result[key], baseline[key], places=9)
                np.testing.assert_allclose(result['adsorption_quantity'], baseline['adsorption_quantity'] * factor)
        self.assertEqual(parsed, original)
        self.assertAlmostEqual(bet.analyze_bet(parsed, .05, .3)['surface_area_m2_g'], baseline['surface_area_m2_g'])

    def test_pressure_window_changes_fit_but_source_report_stays_fixed(self):
        parsed = synthetic_isotherm()
        low = bet.analyze_bet(parsed, .041, .18)
        wide = bet.analyze_bet(parsed, .041, .3)
        self.assertGreater(abs(low['surface_area_m2_g'] - wide['surface_area_m2_g']), .1)
        self.assertLess(low['n_points'], wide['n_points'])
        # Moving bounds within gaps selects the same measured points.
        same = bet.analyze_bet(parsed, .042, .179)
        self.assertEqual(low['n_points'], same['n_points'])
        self.assertAlmostEqual(low['surface_area_m2_g'], same['surface_area_m2_g'])
        self.assertEqual(parsed['source_metrics']['BET surface area'], '24.0 m²/g')

    def test_absolute_uptake_requires_mass_and_is_divided_once(self):
        parsed = synthetic_isotherm()
        reference = bet.analyze_bet(parsed, .05, .3)
        parsed['quantity_basis'] = 'absolute'
        parsed['sample_mass_g'] = None
        for point in parsed['points']:
            point['source_quantity'] *= .1
            point['quantity_cm3_g_stp'] = None
        with self.assertRaisesRegex(ValueError, 'Enter the sample mass'):
            bet.analyze_bet(parsed)
        result = bet.analyze_bet(parsed, .05, .3, sample_mass_g=.1)
        self.assertAlmostEqual(result['surface_area_m2_g'], reference['surface_area_m2_g'])
        self.assertEqual(result['mass_normalization']['quantity_scale_factor'], 10)

    def test_per_gram_data_without_original_mass_cannot_be_renormalized(self):
        parsed = synthetic_isotherm()
        parsed['sample_mass_g'] = None
        result = bet.analyze_bet(parsed, .05, .3)
        self.assertIsNone(result['mass_normalization']['sample_mass_g'])
        with self.assertRaisesRegex(ValueError, 'original sample mass is missing'):
            bet.analyze_bet(parsed, .05, .3, sample_mass_g=.2)

    def test_invalid_mass_constants_and_pressure_are_rejected(self):
        parsed = synthetic_isotherm()
        for field in ('sample_mass_g', 'cross_section_nm2', 'molar_volume_cm3_mol', 'liquid_molar_volume_cm3_mol'):
            for value in (0, -1, 'nan', 'inf', 'invalid'):
                with self.subTest(field=field, value=value), self.assertRaises(ValueError):
                    bet.analyze_bet(parsed, .05, .3, **{field: value})
        for limits in ((.3, .05), (-.1, .3), (.05, 1), ('nan', .3), (.05, 'inf')):
            with self.subTest(limits=limits), self.assertRaises(ValueError):
                bet.analyze_bet(parsed, *limits)

    def test_parser_preserves_original_uptake_and_interprets_mass_units(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, 'synthetic.xlsx')
            for mass in ('0.1000 g', '100 mg', '100mg'):
                with open(path, 'wb') as handle:
                    handle.write(synthetic_workbook(mass))
                parsed = bet.parse_bet_file(path)
                self.assertEqual(parsed['sample_mass_g'], .1)
                self.assertEqual(parsed['quantity_basis'], 'per_mass')
                self.assertEqual(parsed['points'][0]['source_quantity'], parsed['points'][0]['quantity_cm3_g_stp'])
            with open(path, 'wb') as handle:
                handle.write(synthetic_workbook(header='Quantity Adsorbed (cm³ STP)'))
            self.assertEqual(bet.parse_bet_file(path)['quantity_basis'], 'absolute')
        for header in ('Quantity Adsorbed', 'Quantity Adsorbed (mmol/g)',
                       'Quantity Adsorbed (cm³/kg STP)', 'Quantity Adsorbed (cm³ kg⁻¹ STP)'):
            with self.subTest(header=header), self.assertRaises(ValueError):
                bet._quantity_basis(header)

    def test_workbook_and_regenerated_plot_use_corrected_uptake(self):
        parsed = synthetic_isotherm()
        analysis = bet.analyze_bet(parsed, .05, .3, sample_mass_g=.05)
        metadata = {'sample_id': 'Synthetic BET', 'sample_mass_g': .05}
        with tempfile.TemporaryDirectory() as directory:
            context = {'parsed': parsed, 'analysis': analysis, 'metadata': metadata, 'output_dir': directory}
            from modules import characterization_plot
            with patch.object(characterization_plot, 'render_bet_plot', wraps=characterization_plot.render_bet_plot) as renderer:
                result = bet.regenerate_plot(context, {'title': 'Mass correction check'})
            np.testing.assert_allclose(renderer.call_args.args[2], analysis['adsorption_quantity'])
            workbook = load_workbook(result['summary_path'], data_only=False)
            summary = dict(workbook['Summary'].iter_rows(min_col=1, max_col=2, values_only=True))
            self.assertEqual(summary['Sample mass used (g)'], .05)
            self.assertEqual(summary['Original file mass (g)'], .1)
            self.assertAlmostEqual(summary['BET surface area (m²/g)'], analysis['surface_area_m2_g'])
            self.assertAlmostEqual(workbook['Isotherm']['E2'].value, parsed['points'][0]['source_quantity'])
            self.assertAlmostEqual(workbook['Isotherm']['I2'].value, analysis['adsorption_quantity'][0])
            self.assertAlmostEqual(workbook['BET Fit']['B2'].value, analysis['adsorption_quantity'][0])
            self.assertEqual(workbook['Source Metrics']['B2'].value, '24.0 m²/g')
            workbook.close()
            bet.regenerate_plot(context, {'title': 'Second rendering'})
            self.assertAlmostEqual(context['analysis']['surface_area_m2_g'], analysis['surface_area_m2_g'])
            self.assertEqual(parsed['points'][0]['source_quantity'], synthetic_isotherm()['points'][0]['source_quantity'])


class BetApiTests(unittest.TestCase):
    def test_form_changes_recalculate_area_and_validation_returns_400(self):
        import app as server
        client = server.app.test_client()
        with tempfile.TemporaryDirectory() as directory, patch.object(server, 'UPLOAD_DIR', directory):
            def process(identifier, **fields):
                return client.post('/api/process_bet', data={
                    'file': (io.BytesIO(synthetic_workbook()), 'synthetic.xlsx'),
                    'sample_id': identifier, 'output_dir': directory,
                    'p_min': '.05', 'p_max': '.3', **fields,
                }, content_type='multipart/form-data')

            responses = [process('Original'), process('DoubleMass', sample_mass_g='.2'),
                         process('NewWindow', p_max='.18')]
            for response in responses:
                self.assertEqual(response.status_code, 200, response.json)
            original, mass, window = [response.json for response in responses]
            self.assertAlmostEqual(mass['surface_area_m2_g'], original['surface_area_m2_g'] / 2, places=5)
            self.assertNotEqual(window['surface_area_m2_g'], original['surface_area_m2_g'])
            self.assertEqual(mass['mass_normalization']['sample_mass_g'], .2)
            self.assertEqual([r.json['source_surface_area_m2_g'] for r in responses], [24, 24, 24])
            bad = process('InvalidMass', sample_mass_g='0')
            self.assertEqual(bad.status_code, 400)
            self.assertIn('Sample mass', bad.json['error'])
            bad = process('InvalidWindow', p_min='.4', p_max='.1')
            self.assertEqual(bad.status_code, 400)
            self.assertNotIn('trace', bad.json)


if __name__ == '__main__':
    unittest.main()
