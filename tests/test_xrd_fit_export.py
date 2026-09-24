import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook

from modules.xrd import _write_summary_xlsx, run
from modules.xrd.fit_export import fit_parameter_rows
from test_xrd_publication_exports import publication_result


class FitParameterExportTests(unittest.TestCase):
    def test_workbook_is_readable_recipe_without_native_or_raw_dumps(self):
        result = publication_result()
        result['gsas_native_parameters'] = {
            'final_project': {'Phases': 'NATIVE_TREE_MUST_NOT_APPEAR'},
            'input_files': {'cif_text': 'RAW_CIF_MUST_NOT_APPEAR' * 10000},
            'effective_settings': {'n_bg_coeffs': 8},
        }
        result['fit_settings'] = {
            'software': {'toolkit_version': '1.2.2'},
            'project_file': 'sample_xrd_refinement.gpx',
            'source': {'filename': '=scan.xy'},
            'instrument_file': {'filename': 'bench.instprm', 'sha256': 'a' * 64},
            'submitted_parameters': {
                'instrument': 'benchtop_cu', 'wavelength': 1.54056,
                'tt_min': 20, 'tt_max': 90, 'n_bg_coeffs': 'auto',
                'verification_mode': True, 'verify_refine_cell': False,
                'phase_isolation': True, 'verify_use_zero_not_displace': False,
                'verify_refine_x': False, 'verify_fix_y': False,
                'verify_y_nonnegative': True, 'verify_refine_uiso': False,
                'refine_xyz': False, 'size_reporting_mode': 'both', 'scherrer_k': 0.9,
                'phase_options': [{'refine_cell': True, 'refine_size': False,
                    'refine_mustrain': True, 'po_mode': 'fixed', 'po_value': 0.905,
                    'po_axis': [0, 0, 1], 'uniform_cell': False}]},
            'effective_parameters': {'wavelength': 1.540593},
            'interface_settings': {'controls': {'fix_y_value': '7.50',
                'checkboxes': {'xrd-fix-y': False}},
                'phase_options': [{'po_axis': '0 0 0 1', 'po_value': 0.905}]},
            'input_phases': [{'name': 'W2C card', 'formula': 'W2C', 'source': 'mp',
                'mp_id': 'mp-1234', 'spacegroup': 'Pbcn', 'system': 'orthorhombic', 'Z': 4,
                'a': 3.01, 'b': 4.02, 'c': 5.03, 'alpha': 90, 'beta': 90, 'gamma': 90,
                'cif_text': 'RAW_INPUT_CIF_MUST_NOT_APPEAR'}],
            'parsed_input': {'intensity': list(range(100000))},
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = _write_summary_xlsx(result, {'sample_id': '=NOT_A_FORMULA()',
                    'figure_title': '', 'show_figure_title': False}, 'GSAS-II', tmp)
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ['Summary', 'Plot Data', 'Fit Parameters'])
            sheet = workbook['Fit Parameters']
            self.assertEqual([cell.value for cell in sheet[1]], ['Section', 'Setting', 'Value'])
            values = {(r[0].value, r[1].value): r[2].value for r in sheet.iter_rows(min_row=2)}
            self.assertLess(sheet.max_row, 80)
            self.assertEqual(values[('Scan settings', 'Instrument')], 'Benchtop Cu — flat plate (Si 640g)')
            self.assertEqual(values[('Scan settings', 'Background terms')], 'Auto')
            self.assertEqual(values[('Scan settings', 'Background terms used')], 8)
            self.assertEqual(values[('Scan settings', 'Wavelength entered (Å)')], 1.54056)
            self.assertEqual(values[('Scan settings', 'Wavelength used (Å)')], 1.540593)
            self.assertEqual(values[('GSAS controls', 'Fix Y')], 'Unchecked')
            self.assertEqual(values[('GSAS controls', 'Fix Y value (centideg)')], 7.5)
            phase = 'Phase 1: W2C card'
            self.assertEqual(values[(phase, 'Phase card ID')], 'mp-1234')
            self.assertEqual(values[(phase, 'Z (formula units per cell)')], 4)
            self.assertEqual(values[(phase, 'Starting a (Å)')], 3.01)
            self.assertEqual(values[(phase, 'Cell')], 'Checked')
            self.assertEqual(values[(phase, 'Size')], 'Unchecked')
            self.assertEqual(values[(phase, 'Mustrain')], 'Checked')
            self.assertEqual(values[(phase, 'PO axis')], '0 0 0 1')
            self.assertEqual(values[(phase, 'PO value')], 0.905)
            self.assertEqual(values[(phase, 'Uniform cell')], 'Unchecked')
            self.assertEqual(values[('Repeat this fit', 'GSAS-II project')], 'sample_xrd_refinement.gpx')
            self.assertNotIn('MUST_NOT_APPEAR', str(values))
            for row in sheet.iter_rows(min_row=2):
                if isinstance(row[2].value, str):
                    self.assertEqual(row[2].data_type, 's')
                self.assertTrue(row[2].alignment.wrap_text)
            self.assertEqual(sheet.freeze_panes, 'A2')
            self.assertEqual(workbook['Plot Data'].cell(2, 2).value, result['y_obs'][0])
            workbook.close()

    def test_old_results_do_not_invent_initial_controls(self):
        rows = fit_parameter_rows({'phase_results': [{'a': 99}]}, {}, 'GSAS-II')
        self.assertIn('Run the fit again', str(rows))
        self.assertNotIn('Starting a', str(rows))
        self.assertNotIn('Unchecked', str(rows))
        self.assertNotIn('native_state', str(rows))

    def test_phase_order_manual_file_and_missing_settings(self):
        result = {'fit_settings': {'submitted_parameters': {'method': 'gsas2'},
            'input_phases': [
                {'name': 'First', 'cod_id': 'manual', 'source': 'manual', 'cif_filename': 'first.cif'},
                {'name': 'Second', 'cod_id': '1010000', 'source': 'cod'}]}}
        rows = fit_parameter_rows(result, {}, 'GSAS-II')
        values = {(r['Section'], r['Setting']): r['Value'] for r in rows}
        self.assertEqual(values[('Phase 1: First', 'CIF file')], 'first.cif')
        self.assertEqual(values[('Phase 1: First', 'Phase card ID')], 'Manual CIF')
        self.assertEqual(values[('Phase 2: Second', 'Phase card ID')], '1010000')
        self.assertEqual(values[('Phase 1: First', 'Cell')], 'Not recorded')
        self.assertEqual(values[('GSAS controls', 'Quick constrained fit')], 'Not recorded')
        self.assertEqual(values[('GSAS controls', 'Fix Y value (centideg)')], 'Not recorded')

    def test_uniform_cell_is_recorded_for_alternate_w2c_formula(self):
        result = {'fit_settings': {'submitted_parameters': {'phase_options': [{'uniform_cell': True}]},
                                  'input_phases': [{'name': 'Carbide', 'formula': 'CW2'}]}}
        rows = fit_parameter_rows(result, {}, 'GSAS-II')
        self.assertIn({'Section': 'Phase 1: Carbide', 'Setting': 'Uniform cell', 'Value': 'Checked'}, rows)

    def test_generated_cif_is_distinguished_from_reference_card(self):
        result = {'fit_settings': {'submitted_parameters': {'method': 'gsas2'},
            'input_phases': [{'name': 'gamma MoC1-x', 'source': 'generated',
                'cod_id': 'generated:gamma-moc1x-x0.200', 'mp_id': 'mp-2746',
                'generated_cif_model': 'gamma_carbide_fm3m_v1',
                'gamma_vacancy_x': 0.2, 'gamma_c_occupancy': 0.8, 'a': 4.3}]}}
        values = {r['Setting']: r['Value'] for r in fit_parameter_rows(result, {}, 'GSAS-II')}
        self.assertEqual(values['Phase source'], 'Generated carbide CIF')
        self.assertEqual(values['Phase card ID'], 'generated:gamma-moc1x-x0.200')
        self.assertEqual(values['Reference Materials Project card'], 'mp-2746')
        self.assertEqual(values['Carbon vacancy fraction x'], 0.2)
        self.assertEqual(values['Fixed carbon occupancy'], 0.8)
        self.assertEqual(values['Starting a (Å)'], 4.3)

    def test_run_records_input_before_fit_and_exports_native_project(self):
        import base64
        result = publication_result()
        result['_gsas_project_base64'] = base64.b64encode(b'native-project-test').decode()
        phase = {'name': 'W', 'cif_text': 'data_W', 'a': 3.16}
        params = {'method': 'gsas2', 'phases': [phase], 'wavelength': 1.54056,
                  'instprm_file': 'test.instprm', 'legend_location': 'outside right',
                  'tt_min': 20, 'tt_max': 60, 'n_bg_coeffs': 'auto'}
        with tempfile.TemporaryDirectory() as tmp:
            scan = Path(tmp) / 'scan.xye'
            scan.write_text('20 100 2\n30 200 3\n60 100 2\n')
            profile = Path(tmp) / 'test.instprm'
            profile.write_text('#GSAS-II instrument parameter file\nType:PXC\nLam:1.54056\nZero:0\nU:2\nV:-2\n'
                               'W:5\nX:0\nY:0\nSH/L:0.002\n')
            params['instprm_file'] = str(profile)
            with patch('modules.xrd.validate_phases', return_value=[phase]), \
                 patch('modules.xrd.gsasii_backend.is_available', return_value=True), \
                 patch('modules.xrd.gsasii_backend.run_gsas2', return_value=result), \
                 patch('modules.xrd.xrd_plots.make_xrd_plot') as plot:
                exported = run(str(scan), tmp, {'sample_id': 'W'}, params)
            settings = exported['result']['fit_settings']
            self.assertEqual(settings['parsed_input']['sigma'], [2, 3, 2])
            self.assertEqual(settings['effective_parameters']['iteration_limit'], 30)
            self.assertTrue(settings['effective_parameters']['auto_background'])
            phase['a'] = 99
            self.assertEqual(settings['input_phases'][0]['a'], 3.16)
            self.assertEqual(Path(exported['project_path']).read_bytes(), b'native-project-test')
            self.assertNotIn('_gsas_project_base64', exported['result'])
            self.assertEqual(plot.call_args.args[1]['legend_location'], 'outside right')


if __name__ == '__main__':
    unittest.main()
